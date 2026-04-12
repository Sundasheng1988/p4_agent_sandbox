from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.document import Document as _Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.document.normalizer import normalize_document
from app.core.document.parsers.base import BaseDocumentParser
from app.core.document.schema import DocumentBlock, ParsedDocument


_HEADING_RE = re.compile(r"heading\s*([1-9]\d*)?", re.IGNORECASE)


class DocxParser(BaseDocumentParser):
    def parse(self, file_path: Path, file_id: str, filename: str) -> ParsedDocument:
        doc = Document(str(file_path))
        blocks: List[DocumentBlock] = []
        order = 0
        section_stack: List[str] = []

        def next_block_id() -> str:
            return f"{file_id}_b{len(blocks):04d}"

        for item in self._iter_block_items(doc):
            if isinstance(item, Paragraph):
                text = (item.text or "").strip()
                if not text:
                    continue

                level = self._extract_heading_level(item)
                if level is not None:
                    section_stack = section_stack[: max(level - 1, 0)]
                    section_stack.append(text)
                    blocks.append(
                        DocumentBlock(
                            block_id=next_block_id(),
                            block_type="heading",
                            text=text,
                            order=order,
                            level=level,
                            section_path=list(section_stack),
                            meta={"style": item.style.name if item.style else None},
                        )
                    )
                    order += 1
                else:
                    block_type = "list" if self._looks_like_list_paragraph(text) else "paragraph"
                    blocks.append(
                        DocumentBlock(
                            block_id=next_block_id(),
                            block_type=block_type,
                            text=text,
                            order=order,
                            section_path=list(section_stack),
                            meta={"style": item.style.name if item.style else None},
                        )
                    )
                    order += 1

            elif isinstance(item, Table):
                table_text = self._table_to_text(item)
                if not table_text.strip():
                    continue

                blocks.append(
                    DocumentBlock(
                        block_id=next_block_id(),
                        block_type="table",
                        text=table_text,
                        order=order,
                        section_path=list(section_stack),
                        meta={
                            "rows": len(item.rows),
                            "cols": len(item.columns),
                        },
                    )
                )
                order += 1

        parsed = ParsedDocument(
            file_id=file_id,
            filename=filename,
            file_type="docx",
            title=None,
            blocks=blocks,
            meta={},
        )
        return normalize_document(parsed)

    def _iter_block_items(self, doc: _Document):
        """
        优先使用 python-docx 当前文档对象支持的 iter_inner_content()。
        文档说明表明它可以按文档顺序遍历 Paragraph / Table。:contentReference[oaicite:3]{index=3}
        """
        if hasattr(doc, "iter_inner_content"):
            yield from doc.iter_inner_content()
            return

        body = doc.element.body
        for child in body.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, doc)
            elif child.tag.endswith("}tbl"):
                yield Table(child, doc)

    def _extract_heading_level(self, para: Paragraph) -> Optional[int]:
        style_name = para.style.name if para.style else ""
        if not style_name:
            return None

        match = _HEADING_RE.fullmatch(style_name.strip())
        if not match:
            return None

        raw = match.group(1)
        if raw is None:
            return 1
        try:
            return int(raw)
        except ValueError:
            return 1

    def _looks_like_list_paragraph(self, text: str) -> bool:
        stripped = text.strip()
        if stripped.startswith(("- ", "* ", "+ ")):
            return True
        if re.match(r"^\d+[.)]\s+", stripped):
            return True
        return False

    def _table_to_text(self, table: Table) -> str:
        rows: List[str] = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = "\n".join(
                    p.text.strip() for p in cell.paragraphs if (p.text or "").strip()
                ).strip()
                cells.append(cell_text)
            rows.append("| " + " | ".join(cells) + " |")
        return "\n".join(rows).strip()